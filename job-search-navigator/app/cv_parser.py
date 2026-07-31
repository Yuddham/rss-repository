import io
import zipfile
from typing import Any, Optional

SECTION_HEADINGS = {
    "summary": ["summary", "objective", "profile", "about"],
    "experience": ["experience", "work experience", "employment history", "professional experience"],
    "education": ["education", "academic background"],
    "skills": ["skills", "technical skills", "key skills", "core competencies"],
    "certifications": ["certifications", "certificates", "licenses", "licenses & certifications"],
    "contact": ["contact", "contact information"],
}


def _match_heading(line: str) -> Optional[str]:
    clean = line.strip().strip(":").lower()
    if not clean or len(clean) > 40:
        return None
    for key, synonyms in SECTION_HEADINGS.items():
        if clean in synonyms:
            return key
    return None


def _split_sections(raw_text: str) -> dict[str, list[str]]:
    sections: dict[str, list[str]] = {"other": []}
    current = "other"
    for line in (raw_text or "").splitlines():
        stripped = line.strip()
        heading_key = _match_heading(stripped)
        if heading_key:
            current = heading_key
            sections.setdefault(current, [])
            continue
        if stripped:
            sections.setdefault(current, []).append(stripped)
    return sections


def _has_column_gap(xs: list[float], page_width: float) -> bool:
    if len(xs) < 10 or not page_width:
        return False
    mid_low, mid_high = page_width * 0.4, page_width * 0.6
    in_gap = [x for x in xs if mid_low < x < mid_high]
    return (len(in_gap) / len(xs)) < 0.05


def _parse_pdf_fallback(content: bytes, warnings: list[str]) -> dict[str, Any]:
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(content))
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
        return {"raw_text": text, "layout_meta": {"pages": len(reader.pages)}, "warnings": warnings}
    except Exception as exc:  # noqa: BLE001 - last-resort fallback must never raise
        warnings.append(f"Could not parse this PDF at all: {exc}")
        return {"raw_text": "", "layout_meta": {}, "warnings": warnings}


def _parse_pdf(content: bytes) -> dict[str, Any]:
    warnings: list[str] = []
    try:
        import pdfplumber
    except ImportError:
        warnings.append("pdfplumber is not installed; falling back to a simpler parser.")
        return _parse_pdf_fallback(content, warnings)

    text_parts: list[str] = []
    fonts: set[str] = set()
    table_count = 0
    image_count = 0
    multi_column_pages = 0
    pages = 0
    total_chars = 0

    try:
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            pages = len(pdf.pages)
            for page in pdf.pages:
                text = page.extract_text() or ""
                text_parts.append(text)
                total_chars += len(text)

                try:
                    table_count += len(page.find_tables())
                except Exception:  # noqa: BLE001 - table detection is best-effort
                    pass

                image_count += len(page.images)

                for ch in page.chars[:2000]:
                    if ch.get("fontname"):
                        fonts.add(ch["fontname"])

                try:
                    words = page.extract_words()
                    if words and _has_column_gap([w["x0"] for w in words], page.width):
                        multi_column_pages += 1
                except Exception:  # noqa: BLE001 - column heuristic is best-effort
                    pass
    except Exception as exc:  # noqa: BLE001 - defensive: never crash an upload
        warnings.append(f"pdfplumber failed to parse this PDF ({exc}); trying a simpler fallback parser.")
        return _parse_pdf_fallback(content, warnings)

    raw_text = "\n".join(text_parts)
    if len(raw_text.strip()) < 20:
        warnings.append("Very little text could be extracted -- this may be a scanned/image-only PDF.")

    layout_meta = {
        "pages": pages,
        "table_count": table_count,
        "image_count": image_count,
        "textbox_count": 0,
        "font_families": sorted(fonts)[:20],
        "multi_column_pages": multi_column_pages,
        "text_chars_per_page": total_chars / max(pages, 1),
        "header_footer_text": [],
    }
    return {"raw_text": raw_text, "layout_meta": layout_meta, "warnings": warnings}


def _count_textboxes_docx(content: bytes) -> int:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as zf:
            with zf.open("word/document.xml") as f:
                xml = f.read().decode("utf-8", errors="ignore")
        return xml.count("w:txbxContent")
    except Exception:  # noqa: BLE001 - best-effort detection
        return 0


def _parse_docx(content: bytes) -> dict[str, Any]:
    warnings: list[str] = []
    try:
        import docx
    except ImportError:
        warnings.append("python-docx is not installed.")
        return {"raw_text": "", "layout_meta": {}, "warnings": warnings}

    try:
        document = docx.Document(io.BytesIO(content))
    except Exception as exc:  # noqa: BLE001 - defensive: never crash an upload
        warnings.append(f"Could not parse this DOCX file: {exc}")
        return {"raw_text": "", "layout_meta": {}, "warnings": warnings}

    paragraphs = [p.text for p in document.paragraphs]
    raw_text = "\n".join(paragraphs)

    fonts: set[str] = set()
    for p in document.paragraphs:
        for run in p.runs:
            if run.font and run.font.name:
                fonts.add(run.font.name)

    header_footer_text = []
    for section in document.sections:
        try:
            if section.header:
                text = " ".join(p.text for p in section.header.paragraphs if p.text)
                if text.strip():
                    header_footer_text.append(text)
            if section.footer:
                text = " ".join(p.text for p in section.footer.paragraphs if p.text)
                if text.strip():
                    header_footer_text.append(text)
        except Exception:  # noqa: BLE001 - best-effort
            pass

    layout_meta = {
        "pages": None,
        "table_count": len(document.tables),
        "image_count": len(document.inline_shapes),
        "textbox_count": _count_textboxes_docx(content),
        "font_families": sorted(fonts)[:20],
        "multi_column_pages": 0,
        "header_footer_text": header_footer_text,
    }

    if len(raw_text.strip()) < 20:
        warnings.append("Very little text could be extracted from this DOCX file.")

    return {"raw_text": raw_text, "layout_meta": layout_meta, "warnings": warnings}


def parse_cv(content: bytes, file_type: str) -> dict[str, Any]:
    if file_type == "pdf":
        parsed = _parse_pdf(content)
    elif file_type == "docx":
        parsed = _parse_docx(content)
    else:
        return {"raw_text": "", "sections": {}, "layout_meta": {}, "warnings": [f"Unsupported file type: {file_type}"]}

    sections = _split_sections(parsed["raw_text"])
    warnings = list(parsed["warnings"])
    for key in ("experience", "education"):
        if not sections.get(key):
            warnings.append(f"Could not detect a clear '{key}' section heading in this CV.")

    return {
        "raw_text": parsed["raw_text"],
        "sections": sections,
        "layout_meta": parsed.get("layout_meta", {}),
        "warnings": warnings,
    }
