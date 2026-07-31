from pathlib import Path
from typing import Any

from app import utils

SECTION_ORDER = ["contact", "summary", "experience", "education", "skills", "certifications", "other"]


def compare_cv_to_jd(cv_raw_text: str, cv_sections: dict[str, list[str]], jd_text: str, top_n: int) -> dict[str, Any]:
    jd_counts = utils.keyword_frequencies(jd_text)
    jd_top = [kw for kw, _ in jd_counts.most_common(top_n)]

    cv_token_sections: dict[str, set[str]] = {}
    cv_token_counts: dict[str, int] = {}
    cv_original_casing: dict[str, str] = {}

    for section_name, lines in cv_sections.items():
        for line in lines:
            tokens = utils.tokenize(line)
            bigrams = utils.extract_bigrams(tokens)
            for t in tokens + bigrams:
                cv_token_counts[t] = cv_token_counts.get(t, 0) + 1
                cv_token_sections.setdefault(t, set()).add(section_name)
            for word in line.split():
                cleaned = word.strip(".,;:()[]{}\"'").lower()
                if cleaned and cleaned not in cv_original_casing:
                    cv_original_casing[cleaned] = word.strip(".,;:()[]{}\"'")

    matched = []
    missing = []
    for kw in jd_top:
        jd_freq = jd_counts[kw]
        cv_freq = cv_token_counts.get(kw, 0)
        if cv_freq > 0:
            matched.append({
                "keyword": kw,
                "jd_freq": jd_freq,
                "cv_freq": cv_freq,
                "cv_sections": sorted(cv_token_sections.get(kw, [])),
            })
        else:
            missing.append({"keyword": kw, "jd_freq": jd_freq})

    missing.sort(key=lambda m: m["jd_freq"], reverse=True)
    match_score = round((len(matched) / len(jd_top)) * 100, 1) if jd_top else 0.0

    key_skills_line = None
    if matched:
        surface_forms = []
        for m in matched[:10]:
            kw = m["keyword"]
            surface_forms.append(kw if " " in kw else cv_original_casing.get(kw, kw))
        key_skills_line = "Key Skills Match: " + ", ".join(surface_forms)

    return {
        "matched_keywords": matched,
        "missing_keywords": missing,
        "match_score": match_score,
        "key_skills_line": key_skills_line,
    }


def _bullet_relevance(line: str, matched_keyword_set: set[str]) -> int:
    line_tokens = utils.tokenize(line)
    all_terms = set(line_tokens) | set(utils.extract_bigrams(line_tokens))
    return len(all_terms & matched_keyword_set)


def _ordered_sections(cv_sections: dict[str, list[str]]) -> list[tuple[str, list[str]]]:
    ordered = []
    seen = set()
    for section in SECTION_ORDER:
        seen.add(section)
        if cv_sections.get(section):
            ordered.append((section, cv_sections[section]))
    for section, lines in cv_sections.items():
        if section not in seen and lines:
            ordered.append((section, lines))
    return ordered


def generate_tailored_markdown(
    cv_sections: dict[str, list[str]],
    matched_keywords: list[dict],
    missing_keywords: list[dict],
    key_skills_line: str | None,
) -> str:
    """Mechanical only: reorders the CV's own existing lines and appends a
    keyword summary/gap list built strictly from words already in the CV or
    the job description keyword list -- no new sentences are authored."""
    matched_set = {m["keyword"] for m in matched_keywords}
    out: list[str] = []

    if key_skills_line:
        out.append(key_skills_line)
        out.append("")

    for section, lines in _ordered_sections(cv_sections):
        out.append(f"## {section.title()}")
        for line in sorted(lines, key=lambda l: _bullet_relevance(l, matched_set), reverse=True):
            out.append(f"- {line}")
        out.append("")

    out.append("---")
    out.append("## Suggested Gaps To Address (manual edits)")
    if missing_keywords:
        for m in missing_keywords[:20]:
            out.append(f"- {m['keyword']}")
    else:
        out.append("- No significant gaps detected.")

    return "\n".join(out)


def generate_tailored_docx(
    cv_sections: dict[str, list[str]],
    matched_keywords: list[dict],
    missing_keywords: list[dict],
    key_skills_line: str | None,
    output_path: Path,
) -> None:
    import docx

    matched_set = {m["keyword"] for m in matched_keywords}
    document = docx.Document()

    if key_skills_line:
        document.add_paragraph(key_skills_line)

    for section, lines in _ordered_sections(cv_sections):
        document.add_heading(section.title(), level=2)
        for line in sorted(lines, key=lambda l: _bullet_relevance(l, matched_set), reverse=True):
            document.add_paragraph(line, style="List Bullet")

    document.add_heading("Suggested Gaps To Address (manual edits)", level=2)
    if missing_keywords:
        for m in missing_keywords[:20]:
            document.add_paragraph(m["keyword"], style="List Bullet")
    else:
        document.add_paragraph("No significant gaps detected.")

    document.save(str(output_path))
